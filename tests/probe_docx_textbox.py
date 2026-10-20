"""验证两件事（不依赖任何网上说法）：
1. python-docx 读 docx 时，能否拿到「文本框(textbox)」里的文字？
2. 用 zipfile + XML 直读时，会不会因为 mc:Fallback 把文本框内容读两遍？

构造一个模仿 Word 真实输出结构的 docx：mc:AlternateContent 里
  Choice  = w:drawing → wps:txbx → w:txbxContent
  Fallback = w:pict   → v:textbox → w:txbxContent
两份内容相同（Word 就是这么存的）。
"""
import io
import sys
import zipfile
from xml.etree import ElementTree as ET

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
W_T = f"{{{W}}}t"
W_P = f"{{{W}}}p"
MC_FALLBACK = f"{{{MC}}}Fallback"

CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>"""

RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""

DOCUMENT = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W}" xmlns:mc="{MC}"
  xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
  xmlns:wps14="http://schemas.microsoft.com/office/word/2010/wordml"
  xmlns:v="urn:schemas-microsoft-com:vml"
  xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
  xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
  mc:Ignorable="w14 wp14 wps14">
<w:body>
  <w:p><w:r><w:t>第一段正常文字</w:t></w:r></w:p>
  <w:p>
    <w:r><w:t>第二段开头</w:t></w:r>
    <w:r>
      <mc:AlternateContent>
        <mc:Choice Requires="wps14">
          <w:drawing>
            <wp:anchor distT="0" distB="0" distL="0" distR="0">
              <a:graphic>
                <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
                  <wps:wsp>
                    <wps:txbx>
                      <w:txbxContent>
                        <w:p><w:r><w:t>【文本框里的内容】</w:t></w:r></w:p>
                      </w:txbxContent>
                    </wps:txbx>
                  </wps:wsp>
                </a:graphicData>
              </a:graphic>
            </wp:anchor>
          </w:drawing>
        </mc:Choice>
        <mc:Fallback>
          <w:pict>
            <v:shape style="width:100pt;height:30pt">
              <v:textbox>
                <w:txbxContent>
                  <w:p><w:r><w:t>【文本框里的内容】</w:t></w:r></w:p>
                </w:txbxContent>
              </v:textbox>
            </v:shape>
          </w:pict>
        </mc:Fallback>
      </mc:AlternateContent>
    </w:r>
    <w:r><w:t>第二段结尾</w:t></w:r>
  </w:p>
  <w:p><w:r><w:t>第三段正常文字</w:t></w:r></w:p>
</w:body>
</w:document>"""


def build_docx(path: str) -> None:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", CONTENT_TYPES)
        z.writestr("_rels/.rels", RELS)
        z.writestr("word/document.xml", DOCUMENT)


def via_python_docx(path: str):
    import docx

    d = docx.Document(path)
    return [p.text for p in d.paragraphs]


def via_xml(path: str, skip_fallback: bool):
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml")
    root = ET.fromstring(xml)
    out: list[str] = []

    def rec(el):
        if skip_fallback and el.tag == MC_FALLBACK:
            return
        if el.tag == W_P:
            out.append("\n")
        elif el.tag == W_T:
            out.append(el.text or "")
        for ch in el:
            rec(ch)

    rec(root)
    return [s for s in "".join(out).split("\n") if s.strip()]


def main():
    print(f"python = {sys.version.split()[0]}")
    import docx

    print(f"python-docx = {docx.__version__ if hasattr(docx, '__version__') else 'unknown'}")
    path = "probe.docx"
    build_docx(path)

    a = via_python_docx(path)
    b = via_xml(path, skip_fallback=False)
    c = via_xml(path, skip_fallback=True)

    def show(title, rows):
        print(f"\n--- {title} ---")
        for r in rows:
            print(f"  | {r}")
        joined = "".join(rows)
        print(f"  含文本框内容: {'是' if '文本框' in joined else '否'}"
              f" | 出现次数: {joined.count('文本框里的内容')}")

    show("A. python-docx  (doc.paragraphs)", a)
    show("B. XML 直读，不跳过 mc:Fallback", b)
    show("C. XML 直读，跳过 mc:Fallback", c)

    print("\n结论:")
    print(f"  python-docx 丢文本框内容 -> {'确认' if '文本框' not in ''.join(a) else '未复现'}")
    print(f"  不跳 Fallback 会重复     -> {'确认' if ''.join(b).count('文本框里的内容') > 1 else '未复现'}")
    print(f"  跳 Fallback 结果正确     -> {'确认' if ''.join(c).count('文本框里的内容') == 1 else '未复现'}")


if __name__ == "__main__":
    main()
